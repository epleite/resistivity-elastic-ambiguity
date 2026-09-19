#!/usr/bin/env python3
"""Prepare the local Oman GT3A sample table from the official workbook.

The measurements are not redistributed with this repository. Download version
1.2 of the supporting workbook from Hiroshima University repository record
2000060 and pass its path with ``--source``. The source SHA-256 is checked
before any values are extracted.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import math
from pathlib import Path
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "data" / "oman" / "oman_gt3a_sample_level.csv"
SOURCE_SHA256 = "dceaae2ee7e85d233349222447f3c17e50d726b580ee2c629bcc1aed9291ff8b"
SOURCE_RECORD = "https://hiroshima.repo.nii.ac.jp/records/2000060"

HEADERS = (
    "Source row",
    "Core",
    "Section",
    "Depth (m)",
    "Lithology",
    "Sequence",
    "Bulk density",
    "Porosity (%)",
    "Wet Vp mean",
    "Wet Vs mean",
    "R35 mean",
    "Formation factor",
    "log10 R35",
    "log10 F",
    "Depth block",
    "Primary complete",
    "F complete",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def finite_number(value: object) -> float | int | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return value if math.isfinite(float(value)) else None


def log10_positive(value: float | int | None) -> float | None:
    return math.log10(float(value)) if value is not None and value > 0 else None


def extract_rows(source: Path) -> list[dict[str, object]]:
    workbook = load_workbook(source, data_only=True, read_only=True)
    try:
        table_s2 = workbook["Table S2"]
        table_s3 = workbook["Table S3"]

        formation_factor: dict[tuple[object, object, object], float | int | None] = {}
        for row_number in range(4, 98):
            values = [table_s3.cell(row_number, column).value for column in range(1, 42)]
            key = (values[0], values[1], values[4])
            formation_factor[key] = finite_number(values[13])

        rows: list[dict[str, object]] = []
        for row_number in range(4, 98):
            values = [table_s2.cell(row_number, column).value for column in range(1, 63)]
            key = (values[0], values[1], values[4])
            density = finite_number(values[7])
            porosity = finite_number(values[9])
            r35 = finite_number(values[33])
            vp_wet = finite_number(values[54])
            vs_wet = finite_number(values[59])
            factor = formation_factor.get(key)
            log_r35 = log10_positive(r35)
            log_factor = log10_positive(factor)
            rows.append(
                {
                    "source_row": row_number,
                    "core": values[0],
                    "section": values[1],
                    "depth": finite_number(values[4]),
                    "lithology": values[5],
                    "sequence": values[6],
                    "density": density,
                    "porosity": porosity,
                    "vp_wet": vp_wet,
                    "vs_wet": vs_wet,
                    "r35": r35,
                    "factor": factor,
                    "log_r35": log_r35,
                    "log_factor": log_factor,
                }
            )
    finally:
        workbook.close()

    rows.sort(key=lambda row: float(row["depth"]))
    for index, row in enumerate(rows):
        row["block"] = 1 + math.floor(8 * index / 94)
        row["primary_complete"] = all(
            row[key] is not None
            for key in ("porosity", "density", "vp_wet", "vs_wet", "log_r35")
        )
        row["factor_complete"] = all(
            row[key] is not None
            for key in ("porosity", "density", "vp_wet", "vs_wet", "log_factor")
        )
    return rows


def write_rows(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(HEADERS)
        for row in rows:
            writer.writerow(
                (
                    row["source_row"],
                    row["core"],
                    row["section"],
                    row["depth"],
                    row["lithology"],
                    row["sequence"],
                    row["density"],
                    row["porosity"],
                    row["vp_wet"],
                    row["vs_wet"],
                    row["r35"],
                    row["factor"],
                    row["log_r35"],
                    row["log_factor"],
                    row["block"],
                    row["primary_complete"],
                    row["factor_complete"],
                )
            )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source",
        type=Path,
        required=True,
        help=f"version 1.2 workbook downloaded from {SOURCE_RECORD}",
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    source = args.source.expanduser().resolve()
    if not source.is_file():
        parser.error(f"source workbook not found: {source}")
    actual_sha = sha256(source)
    if actual_sha != SOURCE_SHA256:
        parser.error(
            "source workbook SHA-256 does not match the archived version 1.2: "
            f"{actual_sha} != {SOURCE_SHA256}"
        )

    rows = extract_rows(source)
    if len(rows) != 94:
        raise RuntimeError(f"expected 94 physical samples, extracted {len(rows)}")
    write_rows(args.output, rows)
    print(f"Wrote {len(rows)} rows to {args.output}")


if __name__ == "__main__":
    main()
